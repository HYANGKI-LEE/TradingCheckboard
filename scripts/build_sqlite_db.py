"""RawData.xlsx(값만 있는 배포본) -> SQLite DB 변환 (통합 와이드 테이블).

  market(date PK, "<시트>__<시리즈명>" ... 시리즈별 REAL 칼럼)  - 날짜 1행 = 모든 시트 통합
  series(column_name, sheet, col_idx, block_title, field)       - 칼럼 메타
  meta(key, value)                                              - 빌드 정보

사용: python build_sqlite_db.py [--src RawData.xlsx] [--update RawData_master.xlsx] [--dst market.db] [--max-date 2026-10-16]
  --src    : 전체 히스토리를 가진 기준 파일
  --update : (선택) 최신 값을 덮어쓸 파일. 같은 (시트, 칼럼위치, 날짜)에 값이 있으면 이 파일 값을 우선한다.
             히스토리를 줄인(최근 며칠만 남은) 시트가 있어도 기준 파일의 과거 데이터는 보존된다.
"""

import argparse
import datetime as dt
import math
import os
import sqlite3
import sys
import time

_VENDOR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "vendor")
if os.path.isdir(_VENDOR) and _VENDOR not in sys.path:
    sys.path.insert(0, _VENDOR)

import openpyxl
import pandas as pd

DATA_DIR = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "data")


def _to_date(v):
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    if isinstance(v, str):
        try:
            return dt.datetime.strptime(v[:10], "%Y-%m-%d").date()
        except ValueError:
            return None
    return None


def _num(v):
    if isinstance(v, (int, float)) and not isinstance(v, bool) and not (isinstance(v, float) and math.isnan(v)):
        return float(v)
    return None


def _short(sheet_title):
    return sheet_title[5:-1] if sheet_title.startswith("Info(") and sheet_title.endswith(")") else sheet_title


def _headers(row2, row3, n_cols):
    titles, cur = [None] * n_cols, None
    for i in range(n_cols):
        v = row2[i] if i < len(row2) else None
        if v not in (None, "") and not str(v).strip().startswith("단위"):  # "단위: %" 같은 주석 셀은 블록 제목이 아님
            cur = str(v).strip()
        titles[i] = cur
    out = {}
    for i in range(1, n_cols):
        field = row3[i] if i < len(row3) and row3[i] not in (None, "") else ""
        title = titles[i] or f"col{i + 1}"
        out[i] = (title, str(field), f"{title}-{field}" if field else title)
    return out


def read_sheet(ws, max_date):
    rows = ws.iter_rows(values_only=True)
    header = {n: next(rows) for n in range(1, 5)}
    n_cols = max(len(header[2]), len(header[3]))
    hdr = _headers(header[2], header[3], n_cols)
    dates, data = [], []
    for row in rows:
        d = _to_date(row[0]) if row else None
        if d is None or d > max_date:
            continue
        dates.append(d.isoformat())
        data.append([_num(row[i]) if i < len(row) else None for i in range(1, n_cols)])
    df = pd.DataFrame(data, index=pd.Index(dates, name="date"), columns=list(range(1, n_cols)), dtype="float64")
    df = df.loc[:, df.notna().any()]
    df = df[~df.index.duplicated(keep="last")]
    return df, hdr


def build(src, dst, max_date, update=None):
    wb = openpyxl.load_workbook(src, data_only=True, read_only=True)
    wb_up = openpyxl.load_workbook(update, data_only=True, read_only=True) if update else None
    frames, meta_rows, used = [], [], set()
    t0 = time.time()
    for ws in wb.worksheets:
        df, hdr = read_sheet(ws, max_date)
        if wb_up is not None and ws.title in wb_up.sheetnames:
            df_up, hdr_up = read_sheet(wb_up[ws.title], max_date)
            df = df_up.combine_first(df).sort_index()
            hdr = {**hdr, **hdr_up}
        short = _short(ws.title)
        names = []
        for i in df.columns:
            title, field, name = hdr[i]
            col = f"{short}__{name}"
            base, k = col, 2
            while col in used:
                col, k = f"{base}#{k}", k + 1
            used.add(col)
            names.append(col)
            meta_rows.append((col, ws.title, i + 1, title, field))
        df.columns = names
        frames.append(df)
        print(f"  [{ws.title}] 시리즈 {df.shape[1]}개, 행 {df.shape[0]}, 경과 {time.time() - t0:.0f}초", flush=True)
    wide = pd.concat(frames, axis=1).sort_index()
    if os.path.exists(dst):
        os.remove(dst)
    con = sqlite3.connect(dst)
    cols_sql = ", ".join(f'"{c}" REAL' for c in wide.columns)
    con.execute(f'CREATE TABLE market (date TEXT PRIMARY KEY, {cols_sql})')
    ph = ",".join("?" * (wide.shape[1] + 1))
    rows = ((d, *[None if (v != v) else v for v in r]) for d, r in zip(wide.index, wide.itertuples(index=False, name=None)))
    con.executemany(f"INSERT INTO market VALUES ({ph})", rows)
    con.execute("CREATE TABLE series (column_name TEXT PRIMARY KEY, sheet TEXT, col_idx INTEGER, block_title TEXT, field TEXT)")
    con.executemany("INSERT INTO series VALUES (?,?,?,?,?)", meta_rows)
    con.execute("CREATE TABLE meta (key TEXT PRIMARY KEY, value TEXT)")
    con.executemany("INSERT INTO meta VALUES (?,?)", {
        "source_file": os.path.basename(src),
        "built_at": dt.datetime.now().isoformat(timespec="seconds"),
        "max_date_cutoff": max_date.isoformat(),
        "note": "값 0은 엑셀 원본 그대로 저장(커브의 0은 '해당 만기 없음'일 수 있음). 문자 값(조회요청실패 등)은 제외(NULL).",
    }.items())
    con.commit()
    con.execute("VACUUM")
    con.close()
    return wide.shape


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--src", default=os.path.join(DATA_DIR, "RawData.xlsx"))
    ap.add_argument("--dst", default=os.path.join(DATA_DIR, "market.db"))
    ap.add_argument("--update", default=None)
    ap.add_argument("--max-date", default="2099-12-31")
    a = ap.parse_args()
    n_rows, n_cols = build(a.src, a.dst, dt.date.fromisoformat(a.max_date), a.update)
    print(f"완료: {n_rows}일 x {n_cols}시리즈, {a.dst}")
