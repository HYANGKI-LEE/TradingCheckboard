"""
수식이 살아있는 원본(마스터) 엑셀 파일에서 값만 뽑아 가벼운 사본을 만든다.

RawData_master.xlsx는 Infomax IMDH 수식이 들어있는 파일이라 Excel(COM)로 저장할
때마다 공유 수식이 풀리면서 용량이 크게 부풀어오른다(관찰: 40.7MB -> 49.5MB, +22%).
앱(data_loader.py)은 어차피 openpyxl로 값만 읽으므로(data_only=True), 수식은 배포본에
필요 없다 - 그래서 값만 남긴 사본을 만들어 그걸 git에 커밋/배포한다.

read_only/write_only 스트리밍 모드를 써서 대용량 파일도 메모리 적게 쓰고 빠르게 처리한다.
"""

import sys
import openpyxl


def compact(src_path: str, dst_path: str) -> None:
    src_wb = openpyxl.load_workbook(src_path, data_only=True, read_only=True)
    try:
        dst_wb = openpyxl.Workbook(write_only=True)
        for ws in src_wb.worksheets:
            dst_ws = dst_wb.create_sheet(title=ws.title)
            # read_only 모드는 행마다 "그 행에서 마지막으로 값이 있는 칸"까지만 튜플 길이를
            # 반환한다 - 원본은 IMDH 수식이 전체 폭에 깔려있어 모든 행이 시트 전체 너비로
            # 균일하지만(빈 결과라도 수식 자체는 있으니 셀은 존재), data_only로 값만 뽑으면
            # 뒤쪽이 비어있는 행은 짧게 잘려서 나옴. 이 상태로 그대로 append하면 맨 끝 블록
            # (가장 나중에 추가된 상품 - 예: 회사채B, 선물30년) 헤더 행 길이가 짧아져서
            # _find_blocks가 마지막 블록의 끝 칼럼을 잘못 잡아 데이터가 소실된다.
            # -> 시트의 실제 최대 폭(max_column)으로 매 행을 패딩해서 원본과 동일한 폭 유지.
            # openpyxl은 append() 시 "행의 맨 끝쪽에 연속된 None"은 셀 자체를 안 써버린다
            # (파일 크기 절약 최적화). 원본 행이 이미 max_col 길이여도 맨 끝 블록의 나머지
            # 칸들(예: 가장 최근에 추가된 상품이라 첫 칸에만 제목이 있고 나머지가 빈 블록)이
            # 전부 None이면 그 부분이 통째로 잘려서, 다시 읽었을 때 행 길이가 짧아지고
            # _find_blocks가 마지막 블록의 끝 칼럼을 잘못 잡아 데이터가 소실된다.
            # -> 맨 끝 칸이 None이면 공백 문자열로 바꿔서 강제로 셀이 기록되게 한다
            #    (중간의 None들은 그 뒤에 실제 값이 있는 한 정상적으로 보존됨).
            max_col = ws.max_column or 0
            for row in ws.iter_rows(values_only=True):
                if len(row) < max_col:
                    row = row + (None,) * (max_col - len(row))
                if max_col and row and row[-1] is None:
                    row = row[:-1] + (" ",)
                dst_ws.append(row)
        dst_wb.save(dst_path)
    finally:
        src_wb.close()


if __name__ == "__main__":
    if len(sys.argv) != 3:
        print("usage: compact_rawdata.py <src.xlsx> <dst.xlsx>")
        sys.exit(1)
    compact(sys.argv[1], sys.argv[2])
    print(f"압축 완료: {sys.argv[1]} -> {sys.argv[2]}")
